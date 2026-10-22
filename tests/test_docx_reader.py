from pathlib import Path
from docx import Document
from processor.docx_reader import extract_text_from_docx


def test_extract_text_from_docx(tmp_path: Path):
    docx_path = tmp_path / "sample.docx"

    document = Document()
    document.add_heading("Quarterly Revenue", level=1)
    document.add_paragraph("Revenue increased by 12%.")
    document.add_paragraph("Costs decreased by 4%.")

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Quarter"
    table.rows[0].cells[1].text = "Revenue"

    row = table.add_row().cells
    row[0].text = "Q1"
    row[1].text = "$100,000"

    document.save(docx_path)

    result = extract_text_from_docx(docx_path)

    assert result["filename"] == "sample.docx"

    elements = result["elements"]

    assert any(
        item["type"] == "heading"
        and item["text"] == "Quarterly Revenue"
        for item in elements
    )

    assert any(
        item["type"] == "paragraph"
        and item["text"] == "Revenue increased by 12%."
        for item in elements
    )

    table_element = next(
        item for item in elements
        if item["type"] == "table"
    )

    assert table_element["rows"][0] == ["Quarter", "Revenue"]
    assert table_element["rows"][1] == ["Q1", "$100,000"]

def test_extract_text_from_docx_list_items(tmp_path: Path):
    docx_path = tmp_path / "lists.docx"

    document = Document()
    document.add_paragraph(
        "First item",
        style="List Bullet",
    )
    document.add_paragraph(
        "Second item",
        style="List Number",
    )
    document.save(docx_path)

    result = extract_text_from_docx(docx_path)

    list_items = [
        item for item in result["elements"]
        if item["type"] == "list_item"
    ]

    assert len(list_items) == 2
    assert list_items[0]["text"] == "First item"
    assert list_items[1]["text"] == "Second item"