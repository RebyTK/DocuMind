import json
import sys
from pathlib import Path
from docx import Document
import main
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle


def test_docx_pipeline_generates_markdown_and_metadata(
    tmp_path: Path,
    monkeypatch,
):
    documents_dir = tmp_path / "documents"
    documents_dir.mkdir()

    docx_path = documents_dir / "report.docx"

    document = Document()
    document.add_heading("Quarterly Revenue", level=1)
    document.add_paragraph("Revenue increased by 12%.")

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Quarter"
    table.rows[0].cells[1].text = "Revenue"
    document.save(docx_path)

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        ["main.py", str(docx_path)],
    )

    main.main()

    output_dir = tmp_path / "output"
    markdown_path = output_dir / "report_cleaned.md"
    metadata_path = output_dir / "report_metadata.json"

    assert markdown_path.is_file()
    assert metadata_path.is_file()

    markdown = markdown_path.read_text(encoding="utf-8")
    assert "# Quarterly Revenue" in markdown
    assert "Revenue increased by 12%." in markdown
    assert "| Quarter | Revenue |" in markdown

    metadata = json.loads(
        metadata_path.read_text(encoding="utf-8")
    )

    assert metadata["filename"] == "report.docx"
    assert metadata["output_file"] == "report_cleaned.md"
    assert any(
        element["type"] == "table"
        for element in metadata["elements"]
    )


def test_pdf_pipeline_preserves_tables(tmp_path, monkeypatch):
    pdf_path = tmp_path / "report.pdf"

    data = [
        ["Quarter", "Revenue"],
        ["Q1", "$100,000"],
        ["Q2", "$112,000"],
    ]

    document = SimpleDocTemplate(str(pdf_path))

    table = Table(data)
    table.setStyle(
        TableStyle([
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ])
    )

    document.build([table])

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        sys,
        "argv",
        ["main.py", str(pdf_path)],
    )

    main.main()

    output_dir = tmp_path / "output"

    cleaned_path = output_dir / "report_cleaned.txt"
    metadata_path = output_dir / "report_metadata.json"

    assert cleaned_path.is_file()
    assert metadata_path.is_file()

    cleaned_text = cleaned_path.read_text(encoding="utf-8")

    assert "| Quarter | Revenue |" in cleaned_text
    assert "| Q1 | $100,000 |" in cleaned_text

    metadata = json.loads(
        metadata_path.read_text(encoding="utf-8")
    )

    assert "tables" in metadata
    assert len(metadata["tables"]) == 1
    assert metadata["tables"][0]["page"] == 1
    assert metadata["tables"][0]["rows"][1] == ["Q1", "$100,000"]

